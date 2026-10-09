from pathlib import Path

from PIL import Image as PILImage, ImageOps
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "proposta-comercial"
ASSETS = ROOT / "public" / "assets"

BLACK = "070707"
INK = "1B1717"
CREAM = "F5F0E7"
CHALK = "FFFBF4"
GOLD = "D5A33E"
MAGENTA = "C71C86"
MUTED = "6F655C"
LINE = "D8C9AC"
SOFT = "EEE5D5"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LINE, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, name="Arial", size=10.5, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_para(paragraph, before=0, after=8, line=1.12, alignment=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if alignment is not None:
        paragraph.alignment = alignment


def add_text(doc, text, style=None, before=0, after=8, line=1.12, alignment=None):
    p = doc.add_paragraph(style=style)
    set_para(p, before, after, line, alignment)
    run = p.add_run(text)
    return p


def add_body(doc, text, before=0, after=9, line=1.18):
    p = doc.add_paragraph()
    set_para(p, before, after, line)
    set_run_font(p.add_run(text), "Arial", 10.5, INK)
    return p


def add_label(doc, text, after=5):
    p = doc.add_paragraph()
    set_para(p, after=after)
    set_run_font(p.add_run(text.upper()), "Arial", 8.5, GOLD, bold=True)
    return p


def add_heading(doc, text, level=1, before=12, after=8):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_para(p, before, after, 1.0)
    run = p.add_run(text)
    set_run_font(run, "Georgia", 20 if level == 1 else 13.5, BLACK, bold=True)
    return p


def add_bullet(doc, text, after=3):
    p = doc.add_paragraph(style="List Bullet")
    set_para(p, after=after, line=1.08)
    set_run_font(p.add_run(text), "Arial", 10.2, INK)
    return p


def add_accent_line(doc, width=1.18):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(width)
    cell = table.cell(0, 0)
    shade(cell, GOLD)
    cell.height = Inches(0.025)
    cell_margins(cell, 0, 0, 0, 0)
    set_cell_borders(cell, GOLD, "0")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_header(section):
    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    set_para(p, 0, 5, 1.0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run("N  VIDA E BELEZA  |  PITUBA"), "Arial", 8, BLACK, bold=True)
    tab = p.add_run("\tPRESENÇA DIGITAL PREMIUM")
    set_run_font(tab, "Arial", 8, GOLD, bold=True)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.1))
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), LINE)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    set_para(p, 5, 0, 1.0, WD_ALIGN_PARAGRAPH.RIGHT)
    set_run_font(p.add_run("VIDA E BELEZA  ·  PROPOSTA COMERCIAL  "), "Arial", 8, MUTED, bold=True)
    add_page_number(p)


def add_info_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(0.85)
    table.columns[1].width = Inches(5.6)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(0.85)
        cells[1].width = Inches(5.6)
        for cell in cells:
            cell_margins(cell, 55, 0, 55, 100)
            set_cell_borders(cell, CREAM, "0")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[0].paragraphs[0]
        set_para(p, after=0, line=1.0)
        set_run_font(p.add_run(label.upper()), "Arial", 8.5, GOLD, bold=True)
        p = cells[1].paragraphs[0]
        set_para(p, after=0, line=1.0)
        set_run_font(p.add_run(value), "Arial", 9.5, INK)
    return table


def add_comparison_table(doc):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(3.1)
    table.columns[1].width = Inches(3.1)
    headers = ["Instagram", "Landing page"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(3.1)
        shade(cell, BLACK)
        cell_margins(cell, 100, 120, 100, 120)
        set_cell_borders(cell, BLACK, "6")
        p = cell.paragraphs[0]
        set_para(p, after=0)
        set_run_font(p.add_run(text), "Arial", 10, CHALK, bold=True)
    rows = [
        ("Atrai atenção e mantém relacionamento", "Organiza informação e conduz para o contato"),
        ("Depende do formato e do algoritmo", "Oferece um endereço próprio para a clínica"),
        ("Mostra conteúdos em ordem variável", "Apresenta a jornada na ordem planejada"),
        ("É excelente para descoberta social", "Pode apoiar buscas, campanhas e QR Codes"),
    ]
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, text in enumerate(row):
            cell = cells[col_index]
            cell.width = Inches(3.1)
            shade(cell, SOFT if row_index % 2 else CHALK)
            cell_margins(cell, 100, 120, 100, 120)
            set_cell_borders(cell, LINE, "6")
            p = cell.paragraphs[0]
            set_para(p, after=0, line=1.08)
            set_run_font(p.add_run(text), "Arial", 9.6, INK)
    return table


def add_scope_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = [1.35, 3.75, 1.15]
    for col, width in enumerate(widths):
        table.columns[col].width = Inches(width)
    for col, text in enumerate(["Item", "Descrição", "Definição"]):
        cell = table.rows[0].cells[col]
        cell.width = Inches(widths[col])
        shade(cell, BLACK)
        cell_margins(cell, 90, 100, 90, 100)
        set_cell_borders(cell, BLACK, "6")
        p = cell.paragraphs[0]
        set_para(p, after=0)
        set_run_font(p.add_run(text), "Arial", 9, CHALK, bold=True)
    rows = [
        ("Implementação inicial", "Personalização visual, revisão de textos aprovados, substituição das imagens demonstrativas, publicação e validação final.", "A definir"),
        ("Manutenção sob demanda", "Ajustes de conteúdo, campanhas, novas referências visuais e suporte evolutivo quando solicitados.", "A combinar"),
        ("Prazo estimado", "Após o recebimento dos materiais oficiais e a aprovação da direção da clínica.", "A confirmar"),
        ("Pagamento", "Condição comercial a ser combinada entre as partes.", "A combinar"),
    ]
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col, text in enumerate(row):
            cell = cells[col]
            cell.width = Inches(widths[col])
            shade(cell, SOFT if row_index % 2 else CHALK)
            cell_margins(cell, 95, 100, 95, 100)
            set_cell_borders(cell, LINE, "6")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            set_para(p, after=0, line=1.08)
            set_run_font(p.add_run(text), "Arial", 8.8 if col != 0 else 9, INK)
    return table


def add_journey_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = [1.1, 2.2, 2.95]
    for col, width in enumerate(widths):
        table.columns[col].width = Inches(width)
    for col, text in enumerate(["Momento", "Experiência", "Objetivo comercial"]):
        cell = table.rows[0].cells[col]
        shade(cell, BLACK)
        cell_margins(cell, 90, 100, 90, 100)
        set_cell_borders(cell, BLACK, "6")
        p = cell.paragraphs[0]
        set_para(p, after=0)
        set_run_font(p.add_run(text), "Arial", 9, CHALK, bold=True)
    rows = [
        ("Instagram", "Portal de links enxuto", "Direcionar o interesse"),
        ("Primeiro contato", "Landing page premium", "Apresentar a clínica"),
        ("Exploração", "Áreas de cuidado e referências", "Ajudar a pessoa a se identificar"),
        ("Confiança", "Localização, FAQ e linguagem transparente", "Reduzir dúvidas"),
        ("Conversão", "Botão de WhatsApp", "Estimular a avaliação"),
    ]
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col, text in enumerate(row):
            cell = cells[col]
            shade(cell, SOFT if row_index % 2 else CHALK)
            cell_margins(cell, 90, 100, 90, 100)
            set_cell_borders(cell, LINE, "6")
            p = cell.paragraphs[0]
            set_para(p, after=0, line=1.08)
            set_run_font(p.add_run(text), "Arial", 8.8, INK)
    return table


def add_visual_references(doc):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    paths = [ASSETS / "vida-beleza-facial.png", ASSETS / "vida-beleza-corporal.png", ASSETS / "vida-beleza-capilar.png"]
    for col, path in enumerate(paths):
        cell = table.rows[0].cells[col]
        cell.width = Inches(2.05)
        cell_margins(cell, 0, 20, 0, 20)
        set_cell_borders(cell, CREAM, "0")
        p = cell.paragraphs[0]
        set_para(p, after=0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run()
        run.add_picture(str(path), width=Inches(1.95))
    return table


def cover_image_path():
    """Create a proportional cover crop so the editable document never stretches the source."""
    source_path = ASSETS / "vida-beleza-corporal.png"
    crop_path = OUT / "cover-crop.png"
    with PILImage.open(source_path) as source:
        fitted = ImageOps.fit(source.convert("RGB"), (1200, 900), method=PILImage.Resampling.LANCZOS, centering=(0.5, 0.5))
        fitted.save(crop_path, quality=94)
    return crop_path


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    add_page_header(section)
    add_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    for name, size in (("Title", 29), ("Heading 1", 20), ("Heading 2", 13.5)):
        style = styles[name]
        style.font.name = "Georgia"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.font.bold = True

    cover_path = cover_image_path()

    # Page 1 - cover
    add_label(doc, "VIDA E BELEZA  ·  PITUBA  ·  SALVADOR, BA", after=35)
    p = doc.add_paragraph(style="Title")
    set_para(p, after=10, line=0.98)
    set_run_font(p.add_run("Proposta comercial para a Clínica Vida e Beleza"), "Georgia", 28, BLACK, bold=True)
    p = doc.add_paragraph()
    set_para(p, after=22, line=1.08)
    set_run_font(p.add_run("Landing page premium, portal de links e estrutura digital para transformar interesse em conversa."), "Arial", 14, MUTED)
    add_accent_line(doc)
    p = doc.add_paragraph()
    set_para(p, before=8, after=5, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(str(cover_path), width=Inches(6.0), height=Inches(4.5))
    p = doc.add_paragraph()
    set_para(p, after=16, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(p.add_run("Direção visual demonstrativa para uma experiência sofisticada, acolhedora e orientada à conversa."), "Arial", 8.5, MUTED, italic=True)
    add_body(doc, "Esta proposta apresenta uma base de presença digital pensada para a Clínica Vida e Beleza. O objetivo é oferecer uma apresentação mais profissional da clínica, organizar a descoberta das áreas de cuidado e facilitar a solicitação de uma avaliação pelo WhatsApp.", after=17)
    add_info_table(doc, [("Público", "Pessoas que buscam estética avançada, autocuidado e uma conversa clara em Salvador."), ("Status", "Demonstração comercial não oficial")])
    doc.add_page_break()

    # Page 2 - commercial vision
    add_label(doc, "VISÃO COMERCIAL")
    add_heading(doc, "Uma primeira conversa começa antes do WhatsApp", 1, before=4, after=12)
    add_body(doc, "Quando uma pessoa encontra uma clínica pelo Instagram, por uma indicação ou por uma busca na internet, ela precisa entender rapidamente quem é a empresa, como pode começar e qual é o próximo passo.")
    add_body(doc, "A proposta transforma esse interesse inicial em uma conversa mais simples, clara e profissional com a Vida e Beleza. A experiência combina direção visual premium, informação essencial, transparência e chamadas para ação distribuídas ao longo da jornada.", after=18)
    add_heading(doc, "Por que a Vida e Beleza deve investir nesta presença digital", 2, before=8, after=9)
    add_heading(doc, "Uma primeira impressão mais profissional", 2, before=8, after=4)
    add_body(doc, "A página apresenta a clínica com uma linguagem visual sofisticada, acolhedora e coerente com os sinais públicos do Instagram: preto, creme, dourado, acento magenta e símbolo floral demonstrativo.", after=8)
    add_heading(doc, "Mais facilidade para iniciar uma conversa", 2, before=8, after=4)
    add_body(doc, "O visitante encontra o WhatsApp em pontos estratégicos, sem precisar procurar um número ou navegar por várias telas. O caminho principal é objetivo: conhecer, avaliar e começar.", after=8)
    add_heading(doc, "Melhor aproveitamento do tráfego do Instagram", 2, before=8, after=4)
    add_body(doc, "O portal de links organiza os principais caminhos em uma única tela, com acesso para solicitar avaliação, conhecer as áreas, encontrar a localização e abrir o perfil público.", after=8)
    add_heading(doc, "Comunicação mais cuidadosa e confiável", 2, before=8, after=4)
    add_body(doc, "A proposta evita promessas de resultado e informações clínicas sem confirmação. Os materiais demonstrativos podem ser substituídos por imagens autorizadas, textos aprovados e informações oficiais na publicação definitiva.")
    doc.add_page_break()

    # Page 3 - value
    add_label(doc, "VALOR DO INVESTIMENTO")
    add_heading(doc, "Uma landing page amplia o alcance da clínica", 1, before=4, after=12)
    add_body(doc, "O Instagram é importante para relacionamento e descoberta, mas não precisa ser o único ponto de presença digital da Vida e Beleza. Uma landing page cria um endereço próprio para apresentar a clínica e receber pessoas vindas de diferentes canais.")
    add_heading(doc, "Instagram e landing page cumprem papéis diferentes", 2, before=12, after=10)
    add_comparison_table(doc)
    add_heading(doc, "Um destino próprio para quem pesquisa a clínica", 2, before=20, after=8)
    add_body(doc, "Quando alguém pesquisa por clínica de estética em Pituba, estética avançada em Salvador ou uma avaliação para começar um cuidado, uma página própria pode organizar a relevância local e apresentar a clínica em uma sequência planejada.")
    add_body(doc, "A descoberta depende da indexação, da qualidade do conteúdo e de fatores de busca, portanto não representa uma garantia de primeira posição. Ainda assim, cria uma base que pode ser compartilhada em campanhas, anúncios, cartões digitais, mensagens e QR Codes.")
    add_heading(doc, "Mais confiança no momento da decisão", 2, before=13, after=8)
    add_body(doc, "Uma página própria ajuda a validar a pesquisa de quem recebeu uma indicação ou encontrou a clínica no Google. Ela oferece contexto, clareza e informações essenciais antes do contato, tornando a decisão de solicitar uma avaliação mais natural.")
    doc.add_page_break()

    # Page 4 - scope
    add_label(doc, "ESTRUTURA ENTREGUE")
    add_heading(doc, "O que está sendo proposto", 1, before=4, after=10)
    add_body(doc, "Uma experiência digital integrada para apresentar a clínica, organizar a atenção do visitante e tornar o próximo passo mais natural.")
    add_heading(doc, "Landing page premium", 2, before=10, after=5)
    for item in [
        "Hero section com posicionamento, localização e chamada principal.",
        "Apresentação da proposta de cuidado com linguagem acolhedora.",
        "Jornada em três etapas: conhecer, avaliar e começar.",
        "Áreas facial, corporal, capilar e criolipólise sem promessas clínicas.",
        "Galeria de referências visuais demonstrativas e tópicos editáveis.",
        "Seção institucional com localização e link para o Google Maps.",
        "Perguntas frequentes, WhatsApp, Instagram e aviso de demonstração.",
        "Privacidade, cookies e transparência LGPD como parte do escopo mínimo.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Portal de links para Instagram", 2, before=12, after=5)
    for item in [
        "Identidade visual alinhada à landing page.",
        "Solicitar uma avaliação pelo WhatsApp.",
        "Como chegar ao Ed. TK Tower.",
        "Conhecer as áreas de cuidado e abrir o Instagram.",
        "Layout pensado para mobile, favicon próprio e política acessível.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Referências visuais da direção", 2, before=12, after=5)
    add_visual_references(doc)
    p = doc.add_paragraph()
    set_para(p, before=4, after=0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(p.add_run("Referências visuais demonstrativas para composição e direção de conteúdo. Não representam a clínica, equipe ou resultados reais."), "Arial", 8.5, MUTED, italic=True)
    doc.add_page_break()

    # Page 5 - evolution and technical quality
    add_label(doc, "EXPERIÊNCIA E EVOLUÇÃO")
    add_heading(doc, "Uma estrutura preparada para crescer", 1, before=4, after=11)
    add_body(doc, "O projeto foi organizado para que a Vida e Beleza possa atualizar sua comunicação sem precisar reconstruir toda a experiência a cada mudança.")
    add_heading(doc, "Estrutura de conteúdo", 2, before=12, after=5)
    for item in [
        "Conteúdo centralizado para editar textos e referências visuais.",
        "Inclusão de novos tópicos, campanhas e áreas após confirmação.",
        "Edição, remoção e reordenação de referências.",
        "Organização compatível com uma futura conexão a CMS.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Qualidade técnica", 2, before=12, after=5)
    for item in [
        "Vite, React e TypeScript.",
        "Animações GSAP com movimento discreto e suporte a movimento reduzido.",
        "SEO local inicial para Pituba e Salvador.",
        "Acessibilidade semântica, foco de teclado e layout mobile-first.",
        "Testes automatizados para desktop e mobile com Playwright.",
        "Repositório GitHub e publicação preparada na Vercel.",
        "Aviso de cookies e política-base de privacidade com revisão necessária antes da publicação oficial.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Jornada do visitante", 2, before=14, after=8)
    add_journey_table(doc)
    doc.add_page_break()

    # Page 6 - approval and commercial scope
    add_label(doc, "APROVAÇÃO E CONTRATAÇÃO")
    add_heading(doc, "O que a clínica precisará aprovar", 1, before=4, after=10)
    add_body(doc, "A demonstração apresenta uma direção comercial e visual. Para a publicação oficial, o conteúdo precisa ser validado pela Clínica Vida e Beleza.")
    for item in [
        "Logo e identidade visual oficiais, caso substituam a marca demonstrativa.",
        "Imagens institucionais autorizadas.",
        "Lista real de procedimentos e áreas de atendimento.",
        "Informações da equipe, caso sejam apresentadas.",
        "Textos institucionais e perguntas frequentes oficiais.",
        "Uso de depoimentos e casos reais, se houver autorização.",
        "Endereço, horários, formas de contato e domínio oficial.",
        "Controlador, canais, cookies, ferramentas e política final de privacidade.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Escopo comercial sugerido", 2, before=12, after=8)
    add_scope_table(doc)
    add_heading(doc, "Proteção de dados e LGPD", 2, before=15, after=6)
    add_body(doc, "O projeto será desenvolvido com transparência sobre cookies e serviços de terceiros, minimização de dados e ausência de captação desnecessária na landing page. A adequação definitiva dependerá da validação da clínica, da definição do controlador e dos tratamentos efetivamente utilizados na versão oficial.")
    add_heading(doc, "O que diferencia esta proposta", 2, before=12, after=6)
    add_body(doc, "Esta não é apenas uma página com informações da clínica. É uma experiência pensada para organizar a atenção do visitante e tornar o próximo passo mais natural.")
    add_body(doc, "O projeto combina direção visual premium, clareza comercial, conversão pelo WhatsApp, presença local em Salvador, conteúdo visual editável e uma estrutura técnica preparada para evolução.")
    doc.add_page_break()

    # Page 7 - next step
    add_label(doc, "PRÓXIMO PASSO")
    add_heading(doc, "Uma presença digital que começa a conversa", 1, before=4, after=13)
    p = doc.add_paragraph()
    set_para(p, after=16, line=1.1)
    set_run_font(p.add_run("Uma pessoa interessada em cuidado não procura apenas um procedimento.\nEla procura segurança para dar o primeiro passo."), "Georgia", 16, BLACK, bold=True)
    add_body(doc, "Esta proposta ajuda a Vida e Beleza a oferecer essa segurança desde o primeiro contato, com uma presença digital mais clara, mais próxima e preparada para transformar atenção em conversa.", after=12)
    add_body(doc, "Para avançar, a clínica aprova a direção e envia os materiais oficiais. Em seguida, fazemos a personalização final, a revisão da política de privacidade e a publicação.", after=22)
    add_accent_line(doc, width=2.25)
    p = doc.add_paragraph()
    set_para(p, before=26, after=0)
    set_run_font(p.add_run("Demonstração comercial não oficial · valores, prazos, conteúdos e condições sujeitos à aprovação."), "Arial", 8.5, MUTED, italic=True)

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "proposta-comercial-vida-e-beleza.docx"
    doc.save(path)
    print(path)


if __name__ == "__main__":
    build_document()
